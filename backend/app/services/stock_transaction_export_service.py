"""出入库记录导出（盘库用 Excel）。

只负责把已组装好的行数据渲染为多 Sheet 工作簿，不访问数据库：
- 出入库明细：一行一台设备（SN）/一种辅料，单据信息逐行冗余，便于筛选与透视
- 单据汇总：一张单据一行
- 物料进出汇总：按 仓库 + 物料 统计期间进出与净变化
- 盘点表：当前账面库存 + 实盘数（手填）+ 差异（公式）
"""

from __future__ import annotations

import io
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="F97316")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MISMATCH_FILL = PatternFill("solid", fgColor="FFF3BF")
VOIDED_FONT = Font(color="999999", strike=True)
INPUT_FILL = PatternFill("solid", fgColor="EAF4FF")

DETAIL_COLUMNS: Sequence[Tuple[str, str]] = (
    ("document_number", "单据号"),
    ("type_label", "类型"),
    ("source_label", "出库方式"),
    ("operation_time", "时间"),
    ("warehouse_name", "仓库"),
    ("equipment_code", "物料编码"),
    ("equipment_name", "物料名称"),
    ("category_label", "物料类别"),
    ("serial_number", "SN"),
    ("quantity", "数量"),
    ("unit", "单位"),
    ("batch_number", "批次号"),
    ("receiver_name", "领料人/退库人"),
    ("approver_name", "审批人"),
    ("approved_at", "审批时间"),
    ("operator_name", "经办人（仓管）"),
    ("planned_site", "目标站点（计划）"),
    ("actual_site", "实际安装站点"),
    ("actual_cell", "安装小区"),
    ("site_check", "站点核对"),
    ("material_request_no", "申请单号"),
    ("issue_draft_no", "领料单号"),
    ("out_document_number", "关联出库单"),
    ("status_label", "单据状态"),
    ("notes", "备注"),
)

DOCUMENT_COLUMNS: Sequence[Tuple[str, str]] = (
    ("document_number", "单据号"),
    ("type_label", "类型"),
    ("source_label", "出库方式"),
    ("operation_time", "时间"),
    ("warehouse_name", "仓库"),
    ("total_quantity", "总数量"),
    ("main_count", "主设备（台）"),
    ("aux_count", "辅料（件）"),
    ("receiver_name", "领料人/退库人"),
    ("approver_name", "审批人"),
    ("approved_at", "审批时间"),
    ("operator_name", "经办人（仓管）"),
    ("planned_site", "目标站点（计划）"),
    ("material_request_no", "申请单号"),
    ("issue_draft_no", "领料单号"),
    ("out_document_number", "关联出库单"),
    ("status_label", "单据状态"),
    ("notes", "备注"),
)

SUMMARY_COLUMNS: Sequence[Tuple[str, str]] = (
    ("warehouse_name", "仓库"),
    ("equipment_code", "物料编码"),
    ("equipment_name", "物料名称"),
    ("unit", "单位"),
    ("stock_in", "入库"),
    ("stock_out", "出库"),
    ("return_in", "退库收货"),
    ("adjustment", "调整"),
    ("damage", "报损"),
    ("net", "期间净变化"),
)

STOCKTAKE_COLUMNS: Sequence[Tuple[str, str]] = (
    ("warehouse_name", "仓库"),
    ("equipment_code", "物料编码"),
    ("equipment_name", "物料名称"),
    ("category_label", "物料类别"),
    ("unit", "单位"),
    ("book_qty", "账面库存"),
    ("actual_qty", "实盘数（手填）"),
    ("diff", "差异（实盘-账面）"),
    ("remark", "盘点备注"),
)


def _write_sheet(ws, columns: Sequence[Tuple[str, str]], rows: List[dict]) -> None:
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append([_cell_value(row.get(key)) for key, _ in columns])

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    widths = [len(label) * 2 + 2 for _, label in columns]
    for row in rows[:2000]:  # 采样计算列宽，避免大数据量时过慢
        for idx, (key, _) in enumerate(columns):
            text = str(row.get(key) if row.get(key) is not None else "")
            width = sum(2 if ord(ch) > 127 else 1 for ch in text) + 2
            if width > widths[idx]:
                widths[idx] = width
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 8), 48)


def _cell_value(value):
    if value is None:
        return ""
    return value


def build_stock_transactions_workbook(
    *,
    detail_rows: List[dict],
    document_rows: List[dict],
    summary_rows: List[dict],
    stocktake_rows: Optional[List[dict]],
    filter_desc: List[Tuple[str, str]],
) -> bytes:
    wb = Workbook()

    ws_detail = wb.active
    ws_detail.title = "出入库明细"
    _write_sheet(ws_detail, DETAIL_COLUMNS, detail_rows)
    sn_col = [key for key, _ in DETAIL_COLUMNS].index("serial_number") + 1
    for r_idx, row in enumerate(detail_rows, start=2):
        if row.get("site_mismatch"):
            for c_idx in range(1, len(DETAIL_COLUMNS) + 1):
                ws_detail.cell(row=r_idx, column=c_idx).fill = MISMATCH_FILL
        if row.get("instance_is_voided"):
            ws_detail.cell(row=r_idx, column=sn_col).font = VOIDED_FONT

    ws_docs = wb.create_sheet("单据汇总")
    _write_sheet(ws_docs, DOCUMENT_COLUMNS, document_rows)

    ws_summary = wb.create_sheet("物料进出汇总")
    _write_sheet(ws_summary, SUMMARY_COLUMNS, summary_rows)

    if stocktake_rows is not None:
        ws_take = wb.create_sheet("盘点表")
        _write_sheet(ws_take, STOCKTAKE_COLUMNS, stocktake_rows)
        keys = [k for k, _ in STOCKTAKE_COLUMNS]
        book_col = get_column_letter(keys.index("book_qty") + 1)
        actual_col = get_column_letter(keys.index("actual_qty") + 1)
        diff_idx = keys.index("diff") + 1
        actual_idx = keys.index("actual_qty") + 1
        for r_idx in range(2, len(stocktake_rows) + 2):
            ws_take.cell(row=r_idx, column=actual_idx).fill = INPUT_FILL
            ws_take.cell(
                row=r_idx,
                column=diff_idx,
                value=f'=IF({actual_col}{r_idx}="","",{actual_col}{r_idx}-{book_col}{r_idx})',
            )

    ws_info = wb.create_sheet("导出说明")
    ws_info.column_dimensions["A"].width = 18
    ws_info.column_dimensions["B"].width = 80
    for label, value in filter_desc:
        ws_info.append([label, value])
        ws_info.cell(row=ws_info.max_row, column=1).font = Font(bold=True)
    ws_info.append([])
    notes = [
        ("审批人", "出库取物料申请单的审批人；快速出库无审批流程，标记为“快速出库（无审批）”；退库取收货人。"),
        ("目标站点（计划）", "领料/快速出库时选择的目标站点（选填），历史单据为空。"),
        ("实际安装站点", "主设备按 SN 查询最新绑定记录得到；辅料及未安装设备为空。"),
        ("站点核对", "计划与实际站点不一致的行标黄，需重点核查。"),
        ("物料进出汇总", "入库/退库收货/调整为正向，出库/报损为负向；调拨不计入净变化。"),
        ("盘点表", "账面库存为导出时刻的实时数据（仅列出账面不为 0 的物料）；填写“实盘数”后自动计算差异。"),
    ]
    for label, value in notes:
        ws_info.append([label, value])
        ws_info.cell(row=ws_info.max_row, column=1).font = Font(bold=True)
        ws_info.cell(row=ws_info.max_row, column=2).alignment = Alignment(wrap_text=True)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def summarize_material_flow(detail_rows: List[dict]) -> List[dict]:
    """根据明细行按 仓库 + 物料 汇总期间进出。"""
    buckets: Dict[Tuple[str, str], dict] = {}
    for row in detail_rows:
        key = (row.get("warehouse_name") or "", row.get("equipment_code") or "")
        bucket = buckets.get(key)
        if bucket is None:
            bucket = {
                "warehouse_name": row.get("warehouse_name"),
                "equipment_code": row.get("equipment_code"),
                "equipment_name": row.get("equipment_name"),
                "unit": row.get("unit"),
                "stock_in": 0,
                "stock_out": 0,
                "return_in": 0,
                "adjustment": 0,
                "damage": 0,
            }
            buckets[key] = bucket
        flow_key = row.get("flow_key")
        if flow_key in bucket:
            bucket[flow_key] += int(row.get("flow_qty") or 0)

    result = []
    for bucket in buckets.values():
        bucket["net"] = (
            bucket["stock_in"] + bucket["return_in"] + bucket["adjustment"] - bucket["stock_out"] - bucket["damage"]
        )
        result.append(bucket)
    result.sort(key=lambda x: (x.get("warehouse_name") or "", x.get("equipment_code") or ""))
    return result
